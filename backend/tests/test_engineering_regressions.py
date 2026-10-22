import io
from unittest.mock import AsyncMock, MagicMock, patch
from uuid import uuid4

import pytest
from docx import Document as DocxDocument

from app.services.document_processor import chunk_text, count_tokens, extract_text, validate_file


def test_chunks_are_exact_bounded_source_spans():
    text = ('First sentence.\n\nRésumé 中文 <|endoftext|> ' + 'x' * 300 + '\n') * 8
    chunks = chunk_text(text, chunk_size=50, chunk_overlap=10)
    assert chunks[0]['char_start'] == 0
    assert chunks[-1]['char_end'] == len(text)
    for chunk in chunks:
        assert chunk['content'] == text[chunk['char_start']:chunk['char_end']]
        assert chunk['token_count'] == count_tokens(chunk['content']) <= 50
    for before, after in zip(chunks, chunks[1:]):
        assert after['char_start'] <= before['char_end']


def test_pdf_blank_pages_preserve_numbering():
    pages = [MagicMock(), MagicMock(), MagicMock()]
    for page, text in zip(pages, ['Contract begins.', '', 'Warranty lasts 24 months.']):
        page.extract_text.return_value = text
    with patch('app.services.document_processor.PdfReader') as reader:
        reader.return_value.pages = pages
        reader.return_value.is_encrypted = False
        text, count = extract_text(b'fixture', 'pdf')
    assert count == 3
    assert [c['page_number'] for c in chunk_text(text)] == [1, 3]


def test_docx_tables_are_extracted():
    doc = DocxDocument()
    doc.add_paragraph('Contract')
    doc.add_table(rows=1, cols=2).rows[0].cells[1].text = 'Warranty: 24 months'
    stream = io.BytesIO()
    doc.save(stream)
    assert '24 months' in extract_text(stream.getvalue(), 'docx')[0]


@pytest.mark.parametrize('text', [b'', b'  \n\t'])
def test_empty_extraction_rejected(text):
    with pytest.raises(ValueError, match='No readable text'):
        extract_text(text, 'txt')


def test_empty_upload_rejected():
    with pytest.raises(ValueError, match='empty'):
        validate_file('test.txt', 0, 'txt')


def test_chunk_limit_is_explicit(monkeypatch):
    from app.services.document_processor import settings
    monkeypatch.setattr(settings, 'MAX_CHUNKS_PER_DOC', 1)
    with pytest.raises(ValueError, match='chunk limit'):
        chunk_text('Warranty terms. ' * 40, chunk_size=20, chunk_overlap=0)


@pytest.mark.asyncio
async def test_memory_cache_expires_and_does_not_alias(monkeypatch):
    from app.services import cache_service as cache
    monkeypatch.setattr(cache, '_use_redis', False)
    monkeypatch.setattr(cache, '_memory_cache', {})
    monkeypatch.setattr(cache.time, 'monotonic', lambda: 100)
    await cache.set_cached_answer('test', 'question', {'answer': 'original'})
    result = await cache.get_cached_answer('test', 'question')
    result['answer'] = 'changed'
    assert (await cache.get_cached_answer('test', 'question'))['answer'] == 'original'
    monkeypatch.setattr(cache.time, 'monotonic', lambda: 100 + cache.settings.CACHE_TTL_SECONDS)
    assert await cache.get_cached_answer('test', 'question') is None


@pytest.mark.asyncio
async def test_all_primary_analysis_failures_do_not_mark_ready(db_session, monkeypatch):
    from app.api.routes import documents
    from app.models.models import Document, DocumentStatus, User
    from tests.conftest import TestingSessionLocal
    user = User(email='pipeline@example.com', hashed_password='unused')
    db_session.add(user)
    await db_session.flush()
    doc = Document(user_id=user.id, filename='terms.txt', original_filename='terms.txt',
                   file_type='txt', file_size_bytes=20, storage_path='unused', status=DocumentStatus.pending)
    db_session.add(doc)
    await db_session.commit()
    for name in ['summarize_document', 'extract_entities', 'analyze_sentiment']:
        monkeypatch.setattr(documents, name, AsyncMock(side_effect=RuntimeError('provider failed')))
    monkeypatch.setattr(documents, 'extract_keywords', AsyncMock(return_value=[]))
    await documents.process_document_pipeline(doc.id, b'Warranty lasts 24 months.', TestingSessionLocal)
    await db_session.refresh(doc)
    assert doc.status == DocumentStatus.failed


@pytest.mark.asyncio
async def test_followup_retrieval_and_complete_sources(monkeypatch):
    from app.services import ai_service as ai
    retrieval = AsyncMock(return_value=[{'content': 'Warranty lasts 24 months. ' * 30,
        'chunk_index': 0, 'page_number': 3, 'similarity_score': 0.9}])
    generate = AsyncMock(return_value='It lasts 24 months [Source 1].')
    monkeypatch.setattr(ai, 'retrieve_similar_chunks', retrieval)
    monkeypatch.setattr(ai, '_generate_text', generate)
    response = await ai.answer_question(uuid4(), uuid4(), 'How long is it?', uuid4(),
        'terms.pdf', None, history=[{'question': 'What is the warranty?', 'answer': 'A warranty is included.'}])
    assert 'warranty' in retrieval.call_args.kwargs['query']
    assert len(response.sources[0].content) > 320
    assert response.sources[0].page_number == 3


@pytest.mark.asyncio
async def test_invalid_embedding_rejected(monkeypatch):
    from app.services import rag_service as rag
    response = MagicMock(data=[{'embedding': [0.0, 0.0]}])
    monkeypatch.setattr(rag.litellm, 'aembedding', AsyncMock(return_value=response))
    with pytest.raises(ValueError, match='invalid'):
        await rag.embed_texts(['text'])


def test_password_bytes_and_blank_questions_rejected():
    from pydantic import ValidationError

    from app.schemas.schemas import QueryRequest, RegisterRequest
    with pytest.raises(ValidationError):
        RegisterRequest(email='test@example.com', password='Aa1'+'é'*36)
    with pytest.raises(ValidationError):
        QueryRequest(document_id=uuid4(), question='     ')


@pytest.mark.asyncio
async def test_json_auth_failure_not_retried_as_plain(monkeypatch):
    import litellm

    from app.services import ai_service as ai
    completion = AsyncMock(side_effect=litellm.AuthenticationError('invalid',model='test',llm_provider='gemini'))
    monkeypatch.setattr(ai.litellm,'acompletion',completion)
    with pytest.raises(litellm.AuthenticationError):
        await ai._generate_json('fixture')
    assert completion.call_count == 1


@pytest.mark.asyncio
async def test_daily_quota_is_not_retried():
    import litellm

    from app.services.llm_limiter import call_with_limits
    request = AsyncMock(side_effect=litellm.RateLimitError(
        'GenerateRequestsPerDayPerProjectPerModel-FreeTier retryDelay: 29350s',
        model='test', llm_provider='gemini'))
    with pytest.raises(litellm.RateLimitError):
        await call_with_limits(request)
    assert request.call_count == 1


@pytest.mark.asyncio
async def test_pipeline_timeout_is_persisted(db_session, monkeypatch):
    import asyncio

    from app.api.routes import documents
    from app.models.models import Document, DocumentStatus, User
    from tests.conftest import TestingSessionLocal
    user = User(email='timeout@example.com', hashed_password='unused')
    db_session.add(user)
    await db_session.flush()
    doc = Document(user_id=user.id, filename='terms.txt', original_filename='terms.txt',
                   file_type='txt', file_size_bytes=20, storage_path='unused', status=DocumentStatus.pending)
    db_session.add(doc)
    await db_session.commit()
    async def slow(*args, **kwargs):
        await asyncio.sleep(10)
    monkeypatch.setattr(documents.settings, 'DOCUMENT_PROCESSING_TIMEOUT_SECONDS', 0.05)
    monkeypatch.setattr(documents, 'summarize_document', slow)
    await documents.process_document_pipeline(doc.id, b'Warranty lasts 24 months.', TestingSessionLocal)
    await db_session.refresh(doc)
    assert doc.status == DocumentStatus.failed
    assert 'timed out' in doc.error_message


@pytest.mark.asyncio
async def test_failed_document_retry_preserves_identity(async_client):
    from app.api.routes import documents
    auth = await async_client.post('/api/v1/auth/register', json={
        'email': 'retry@example.com', 'password': 'SecurePassword123!',
    })
    headers = {'Authorization': 'Bearer ' + auth.json()['access_token']}
    with patch.object(documents, 'summarize_document', AsyncMock(side_effect=RuntimeError('provider down'))), \
         patch.object(documents, 'extract_entities', AsyncMock(side_effect=RuntimeError('provider down'))), \
         patch.object(documents, 'analyze_sentiment', AsyncMock(side_effect=RuntimeError('provider down'))):
        upload = await async_client.post('/api/v1/documents/upload', headers=headers,
            files={'file': ('retry.txt', b'Warranty lasts 24 months.', 'text/plain')})
    doc_id = upload.json()['id']
    assert (await async_client.get('/api/v1/documents/' + doc_id, headers=headers)).json()['status'] == 'failed'
    retry = await async_client.post('/api/v1/documents/' + doc_id + '/retry', headers=headers)
    assert retry.status_code == 202
    assert retry.json()['id'] == doc_id
    assert (await async_client.get('/api/v1/documents/' + doc_id, headers=headers)).json()['status'] == 'ready'
    assert (await async_client.get('/api/v1/documents/', headers=headers)).json()['total'] == 1
    assert (await async_client.post('/api/v1/documents/' + doc_id + '/retry', headers=headers)).status_code == 409
    other = await async_client.post('/api/v1/auth/register', json={'email': 'retry-outsider@example.com', 'password': 'SecurePassword123!'})
    outsider = {'Authorization': 'Bearer ' + other.json()['access_token']}
    assert (await async_client.post('/api/v1/documents/' + doc_id + '/retry', headers=outsider)).status_code == 404


@pytest.mark.asyncio
async def test_demo_cannot_retry(async_client):
    auth = await async_client.post('/api/v1/auth/demo')
    response = await async_client.post(f'/api/v1/documents/{uuid4()}/retry',
        headers={'Authorization': 'Bearer ' + auth.json()['access_token']})
    assert response.status_code == 403


@pytest.mark.asyncio
async def test_reindex_replaces_stale_passages(db_session):
    from sqlalchemy import select

    from app.models.models import Document, DocumentChunk, User
    from app.services.rag_service import build_document_index
    user = User(email='reindex@example.com', hashed_password='unused')
    db_session.add(user)
    await db_session.flush()
    doc = Document(user_id=user.id, filename='test.txt', original_filename='test.txt',
        file_type='txt', file_size_bytes=20, storage_path='unused')
    db_session.add(doc)
    await db_session.flush()
    db_session.add_all([DocumentChunk(document_id=doc.id, chunk_index=i, content='old passage', token_count=2) for i in range(3)])
    await db_session.commit()
    await build_document_index(db_session, doc.id, [{'chunk_index': 0, 'content': 'current passage', 'token_count': 2, 'page_number': 4}])
    chunks = (await db_session.execute(select(DocumentChunk).where(DocumentChunk.document_id == doc.id))).scalars().all()
    assert len(chunks) == 1 and chunks[0].content == 'current passage' and chunks[0].page_number == 4


@pytest.mark.asyncio
async def test_interrupted_job_is_recoverable_and_owner_scoped(async_client, db_session):
    from datetime import datetime, timedelta, timezone

    from sqlalchemy import select

    from app.models.models import Document, DocumentStatus, User
    auth = await async_client.post('/api/v1/auth/register', json={'email': 'recover@example.com', 'password': 'SecurePassword123!'})
    owner = (await db_session.execute(select(User).where(User.email == 'recover@example.com'))).scalar_one()
    outsider = User(email='recover-other@example.com', hashed_password='unused')
    db_session.add(outsider)
    await db_session.flush()
    def stalled(user_id):
        return Document(user_id=user_id, filename='test.txt', original_filename='test.txt',
            file_type='txt', file_size_bytes=20, storage_path='unused', status=DocumentStatus.processing,
            updated_at=datetime.now(timezone.utc)-timedelta(hours=1))
    own_doc, other_doc = stalled(owner.id), stalled(outsider.id)
    db_session.add_all([own_doc, other_doc])
    await db_session.commit()
    response = await async_client.get('/api/v1/documents/', headers={'Authorization': 'Bearer '+auth.json()['access_token']})
    assert response.json()['items'][0]['status'] == 'failed'
    assert 'interrupted' in response.json()['items'][0]['error_message']
    await db_session.refresh(other_doc)
    assert other_doc.status == DocumentStatus.processing


@pytest.mark.asyncio
async def test_combined_analysis_uses_one_validated_generation(mock_ai_and_gemini):
    from app.services.ai_service import analyze_document
    result = await analyze_document([{'content': 'Warranty lasts 24 months.'}], 'terms.txt')
    assert result.summary == 'Test Summary'
    assert result.sentiment.label == 'Positive'
    assert mock_ai_and_gemini['acompletion'].call_count == 1


@pytest.mark.asyncio
async def test_combined_ingestion_is_ready_with_all_analysis_panels(async_client, monkeypatch):
    from app.api.routes import documents
    monkeypatch.setattr(documents.settings, 'ANALYSIS_MODE', 'combined')
    auth = await async_client.post('/api/v1/auth/register', json={'email': 'combined@example.com', 'password': 'SecurePassword123!'})
    headers = {'Authorization': 'Bearer '+auth.json()['access_token']}
    uploaded = await async_client.post('/api/v1/documents/upload', headers=headers,
        files={'file': ('terms.txt', b'Warranty lasts 24 months.', 'text/plain')})
    response = await async_client.get('/api/v1/documents/'+uploaded.json()['id']+'/analysis', headers=headers)
    assert response.status_code == 200
    assert response.json()['summary'] and response.json()['entities'] and response.json()['sentiment']


@pytest.mark.asyncio
async def test_provider_quota_failure_explains_recovery(async_client, monkeypatch):
    import litellm

    from app.api.routes import documents
    monkeypatch.setattr(documents.settings, 'ANALYSIS_MODE', 'combined')
    monkeypatch.setattr(documents, 'analyze_document', AsyncMock(side_effect=litellm.RateLimitError('private quota details', model='test', llm_provider='gemini')))
    auth = await async_client.post('/api/v1/auth/register', json={'email': 'quota-retry@example.com', 'password': 'SecurePassword123!'})
    headers = {'Authorization': 'Bearer '+auth.json()['access_token']}
    upload = await async_client.post('/api/v1/documents/upload', headers=headers,
        files={'file': ('terms.txt', b'Warranty lasts 24 months.', 'text/plain')})
    metadata = await async_client.get('/api/v1/documents/'+upload.json()['id'], headers=headers)
    assert metadata.json()['status'] == 'failed'
    assert 'quota' in metadata.json()['error_message']
    assert 'private quota details' not in metadata.text


def test_local_original_is_private_and_not_overwritten(tmp_path):
    import stat

    from app.api.routes.documents import _write_local_original
    path = tmp_path / 'private' / 'terms.txt'
    _write_local_original(path, b'Private document')
    assert stat.S_IMODE(path.stat().st_mode) == 0o600
    with pytest.raises(FileExistsError):
        _write_local_original(path, b'Unexpected replacement')
    assert path.read_bytes() == b'Private document'
